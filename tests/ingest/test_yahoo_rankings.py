"""yahoo_rankings: header-row location, xlsx parsing (banner rows, accents, status flags, ADP
fallback), and end-to-end resolution onto player_id via the shared external_rankings contract.
"""
from __future__ import annotations

import pandas as pd
import pytest
from openpyxl import Workbook

from src.ingest import yahoo_rankings as yr

# --------------------------------------------------------------------------- fixture builder

BANNER_ROWS = [
    ("Yahoo Fantasy Basketball Draft Analysis", None, None, None, None, None, None, None, None),
    ("Standard scoring ... ALL positions ... Ordered as shown by Yahoo", None, None, None, None, None, None, None, None),
    ("Captured 2026-09-28 ... 8 players ... Verification details on Source & Notes", None, None, None, None, None, None, None, None),
    (None, None, None, None, None, None, None, None, None),
]

HEADER_ROW = ("Yahoo Display Order", "Player", "Team", "Eligible Positions", "Status", "Rank",
             "% Drafted", "Preseason ADP", "All Drafts ADP")


def make_workbook(path, data_rows, source_notes=None):
    wb = Workbook()
    ws = wb.active
    ws.title = yr.PLAYERS_SHEET
    for row in BANNER_ROWS:
        ws.append(row)
    ws.append(HEADER_ROW)
    for row in data_rows:
        ws.append(row)

    notes_ws = wb.create_sheet(yr.SOURCE_NOTES_SHEET)
    for row in (source_notes or [("Captured", "2026-09-28"), ("Players", "8")]):
        notes_ws.append(row)

    wb.save(path)
    return path


DATA_ROWS = [
    (1, "Victor Wembanyama", "SAS", "C", None, 1, 1, 1.6, 1.6),
    (2, "Nikola Jokić", "DEN", "C", None, 2, 1, 1.9, 1.9),
    (3, "Luka Dončić", "LAL", "PG,SG", None, 3, 1, 3.6, 3.6),
    (4, "Someone Preseason Only", "BOS", "SF", None, 4, 1, 5.5, None),
    (5, "Someone No Adp At All", "NYK", "PF", None, 5, 1, None, None),
    (6, "Giannis Antetokounmpo", "MIA", "PF,C", "P", 9, 1, 7.8, 7.8),
    (7, "Someone Questionable", "PHX", "SG", "Q", 20, 1, 25.0, 25.0),
    (8, "Someone Out", "CHI", "C", "O", 30, 1, 40.0, 40.0),
]


@pytest.fixture
def workbook_path(tmp_path):
    path = tmp_path / "yahoo_draft_analysis.xlsx"
    return make_workbook(path, DATA_ROWS)


# --------------------------------------------------------------------------- parse_yahoo_xlsx

def test_parse_returns_expected_columns(workbook_path):
    df = yr.parse_yahoo_xlsx(workbook_path)
    assert list(df.columns) == ["source_id", "source_name_raw", "name_parsed", "team", "positions",
                                "status_tag", "ext_rank", "adp", "ecr_vs_adp"]
    assert len(df) == len(DATA_ROWS)


def test_parse_preserves_row_order_by_display_order(workbook_path):
    df = yr.parse_yahoo_xlsx(workbook_path)
    assert df["source_id"].tolist() == [str(i) for i in range(1, 9)]
    assert df["ext_rank"].tolist() == [1, 2, 3, 4, 5, 9, 20, 30]


def test_parse_preserves_accented_names_exactly(workbook_path):
    df = yr.parse_yahoo_xlsx(workbook_path)
    names = dict(zip(df["source_id"], df["source_name_raw"]))
    assert names["2"] == "Nikola Jokić"
    assert names["3"] == "Luka Dončić"
    # name_parsed is the same verbatim value, per the module's contract
    parsed = dict(zip(df["source_id"], df["name_parsed"]))
    assert parsed["2"] == "Nikola Jokić"
    assert parsed["3"] == "Luka Dončić"


def test_parse_status_tag_passthrough_and_null_handling(workbook_path):
    df = yr.parse_yahoo_xlsx(workbook_path)
    status = dict(zip(df["source_id"], df["status_tag"]))
    assert pd.isna(status["1"])         # blank status -> real null (pandas NaN), not "None"/"NA" string
    assert status["6"] == "P"
    assert status["7"] == "Q"
    assert status["8"] == "O"


def test_parse_adp_falls_back_to_preseason_when_all_drafts_blank(workbook_path):
    df = yr.parse_yahoo_xlsx(workbook_path)
    row = df[df["source_id"] == "4"].iloc[0]
    assert row["adp"] == pytest.approx(5.5)


def test_parse_adp_is_none_when_both_blank(workbook_path):
    df = yr.parse_yahoo_xlsx(workbook_path)
    row = df[df["source_id"] == "5"].iloc[0]
    assert pd.isna(row["adp"])          # never 0, never a string -- a real null


def test_parse_multi_position_string_verbatim(workbook_path):
    df = yr.parse_yahoo_xlsx(workbook_path)
    row = df[df["source_id"] == "3"].iloc[0]
    assert row["positions"] == "PG,SG"


def test_parse_ecr_vs_adp_always_none(workbook_path):
    df = yr.parse_yahoo_xlsx(workbook_path)
    assert df["ecr_vs_adp"].isna().all() or (df["ecr_vs_adp"] == None).all()  # noqa: E711


def test_parse_missing_players_sheet_raises(tmp_path):
    path = tmp_path / "no_players_sheet.xlsx"
    wb = Workbook()
    wb.active.title = "SomethingElse"
    wb.save(path)
    with pytest.raises(yr.YahooRankingsError):
        yr.parse_yahoo_xlsx(path)


def test_parse_missing_header_row_raises(tmp_path):
    path = tmp_path / "no_header.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = yr.PLAYERS_SHEET
    for row in BANNER_ROWS:
        ws.append(row)
    # no header row, no data
    wb.save(path)
    with pytest.raises(yr.YahooRankingsError):
        yr.parse_yahoo_xlsx(path)


# --------------------------------------------------------------------------- read_source_notes

def test_read_source_notes_returns_item_details(workbook_path):
    notes = yr.read_source_notes(workbook_path)
    assert list(notes.columns) == ["Item", "Details"]
    assert len(notes) >= 1


def test_read_source_notes_missing_sheet_returns_empty_frame(tmp_path):
    path = tmp_path / "no_notes.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = yr.PLAYERS_SHEET
    for row in BANNER_ROWS:
        ws.append(row)
    ws.append(HEADER_ROW)
    ws.append(DATA_ROWS[0])
    wb.save(path)
    notes = yr.read_source_notes(path)
    assert list(notes.columns) == ["Item", "Details"]
    assert len(notes) == 0


# --------------------------------------------------------------------------- load_yahoo_rankings

def make_players(rows):
    """rows: list of (player_id, player_name, from_year, to_year)."""
    return pd.DataFrame(rows, columns=["player_id", "player_name", "from_year", "to_year"])


def test_load_yahoo_rankings_exact_match(workbook_path):
    players = make_players([
        (100, "Nikola Jokić", None, None),  # exact same accented spelling
    ])
    result = yr.load_yahoo_rankings(workbook_path, players)
    row = result.frame[result.frame["source_id"] == "2"].iloc[0]
    assert bool(row["matched"]) is True
    assert row["player_id"] == 100
    assert row["match_method"] == "exact"


def test_load_yahoo_rankings_accent_stripped_still_counts_as_exact(workbook_path):
    """normalize_name accent-strips BOTH sides to the same key, so an accented Yahoo name against an
    unaccented players-table name hits match_players' by_norm dict directly -- that's the "exact"
    method (by_norm), not "normalized" (which is reserved for the manual alias table)."""
    players = make_players([
        (101, "Nikola Jokic", None, None),  # unaccented in the players table
    ])
    result = yr.load_yahoo_rankings(workbook_path, players)
    row = result.frame[result.frame["source_id"] == "2"].iloc[0]
    assert bool(row["matched"]) is True
    assert row["player_id"] == 101
    assert row["match_method"] == "exact"


def test_load_yahoo_rankings_unmatched_name_kept_as_row(workbook_path):
    players = make_players([
        (102, "Someone Else Entirely", None, None),
    ])
    result = yr.load_yahoo_rankings(workbook_path, players)
    # "Someone Preseason Only" (source_id "4") has no counterpart in `players` at all
    row = result.frame[result.frame["source_id"] == "4"].iloc[0]
    assert bool(row["matched"]) is False
    assert pd.isna(row["player_id"])
    assert row["match_method"] == "unmatched"
    # still present as a row, never silently dropped
    assert len(result.frame) == len(DATA_ROWS)


def test_load_yahoo_rankings_full_frame_shape(workbook_path):
    from src.ingest.external_rankings import RESOLVED_COLUMNS

    players = make_players([(100, "Nikola Jokić", None, None)])
    result = yr.load_yahoo_rankings(workbook_path, players)
    assert list(result.frame.columns) == list(RESOLVED_COLUMNS)
    assert (result.frame["source"] == "yahoo").all()
    assert len(result.frame) == len(DATA_ROWS)


# --------------------------------------------------------------------------- CLI

def test_build_parser_defaults():
    parser = yr.build_parser()
    args = parser.parse_args([])
    assert args.path is None
    assert args.data_dir is None


def test_default_path_uses_env_var(monkeypatch, tmp_path):
    monkeypatch.setenv(yr.DEFAULT_PATH_ENV, str(tmp_path / "custom.xlsx"))
    assert yr.default_path() == tmp_path / "custom.xlsx"


def test_main_reports_failure_cleanly(tmp_path, capsys, monkeypatch):
    monkeypatch.setenv("NBA_DATA_DIR", str(tmp_path / "data"))
    rc = yr.main(["--path", str(tmp_path / "does_not_exist.xlsx"), "--data-dir", str(tmp_path / "data")])
    assert rc == 1
    captured = capsys.readouterr()
    assert "yahoo rankings ingest failed" in captured.err


def _bad_workbook(tmp_path, row):
    return make_workbook(tmp_path / "bad.xlsx", [DATA_ROWS[0], row])


@pytest.mark.parametrize("row,msg", [
    ((2, "X", "DEN", "C", None, None, 1, 1.9, 1.9), "blank Rank"),
    ((2, "X", "DEN", "C", None, "abc", 1, 1.9, 1.9), "non-numeric Rank"),
    ((2, "X", "DEN", "C", None, 2, 1, "abc", "abc"), "non-numeric ADP"),
    (("abc", "X", "DEN", "C", None, 2, 1, 1.9, 1.9), "non-numeric Yahoo Display Order"),
    ((1, "Dup", "DEN", "C", None, 2, 1, 1.9, 1.9), "duplicate Yahoo Display Order"),
])
def test_bad_cells_raise_a_clean_error(tmp_path, row, msg):
    with pytest.raises(yr.YahooRankingsError, match=msg):
        yr.parse_yahoo_xlsx(_bad_workbook(tmp_path, row))


def test_a_corrupt_xlsx_raises_a_clean_error(tmp_path):
    p = tmp_path / "broken.xlsx"
    p.write_bytes(b"this is not a zip file")
    with pytest.raises(yr.YahooRankingsError, match="could not open"):
        yr.parse_yahoo_xlsx(p)
